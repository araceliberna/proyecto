"""Simula las formulas INDEX/MATCH/COUNTIF de la hoja Actualizado y compara con el filtrado real."""
import re, openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

F = 'Diccionario_Supply_Chain_Ripley_Completo_1.xlsx'
wb = openpyxl.load_workbook(F)
dic = wb['Diccionario - Proyecto E2E ']
ws  = wb['Actualizado']
FIRST, LAST, R0, RN = 2, 58, 8, 64

data = {r: [dic.cell(r, c).value for c in range(1, 15)] for r in range(FIRST, LAST+1)}

def sim(level):
    # columna auxiliar P: rank acumulado
    helper = []
    for r in range(R0, RN+1):
        srow = FIRST + (r - R0)
        f = ws.cell(r, 16).value
        assert f'$M{srow}=' in f, (r, f)                    # la fila fuente referida es la correcta
        assert f'$M${FIRST}:$M{srow}' in f, (r, f)          # el rango acumulado es el correcto
        if dic.cell(srow, 13).value == level:
            helper.append(sum(1 for x in range(FIRST, srow+1) if dic.cell(x,13).value == level))
        else:
            helper.append('')
    out = []
    for r in range(R0, RN+1):
        n = r - R0 + 1                                      # ROWS($A$8:$A r)
        pos = helper.index(n) + 1 if n in helper else None   # MATCH exacto
        out.append(None if pos is None else data[FIRST + pos - 1])
    return out

for level in ['Estratégico', 'Táctico', 'Operativo']:
    got = [row for row in sim(level) if row]
    exp = [data[r] for r in range(FIRST, LAST+1) if dic.cell(r,13).value == level]
    print(f'{level:12} simulado={len(got):2}  esperado={len(exp):2}  identico={got == exp}')
    if got != exp:
        print('   DIFERENCIA'); break

# verifica que cada formula de la tabla apunte a su propia columna
bad = []
for c in range(1, 15):
    L = get_column_letter(c)
    for r in range(R0, RN+1):
        f = ws.cell(r, c).value
        m = re.search(r"!\$?([A-Z]+)\$2:", f)
        if m.group(1) != L: bad.append((r, c, f))
        if f'ROWS($A$8:$A{r})' not in f: bad.append((r, c, 'ROWS mal', f))
print('formulas con columna/ROWS incorrecta:', len(bad))
print('B5 =', ws['B5'].value, '| E5 =', ws['E5'].value)
print('A8 =', ws['A8'].value)
print('P8 =', ws['P8'].value)
