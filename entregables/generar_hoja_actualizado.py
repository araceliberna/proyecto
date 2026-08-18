import sys, shutil
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

MODE = sys.argv[1] if len(sys.argv) > 1 else 'helper'   # 'aggregate' | 'helper'

SRC = '/root/.claude/uploads/7bd8d737-48db-5837-beae-ecf7b754ea56/b46dfb99-Diccionario_Supply_Chain_Ripley_Completo_1.xlsx'
OUT = '/tmp/claude-0/-home-user-proyecto/7bd8d737-48db-5837-beae-ecf7b754ea56/scratchpad/Diccionario_Supply_Chain_Ripley_Completo_1.xlsx'
shutil.copy(SRC, OUT)

DIC = 'Diccionario - Proyecto E2E '     # ojo: termina en espacio
Q = f"'{DIC}'"
FIRST, LAST = 2, 58                     # filas de datos en el diccionario
N = LAST - FIRST + 1                    # 57
HDR_ROW = 7
R0 = HDR_ROW + 1                        # 8  -> primera fila de resultados
RN = R0 + N - 1                         # 64
LVL = '$B$5'                            # celda de selección de nivel

wb = openpyxl.load_workbook(OUT)
src = wb[DIC]
if 'Actualizado' in wb.sheetnames:
    del wb['Actualizado']
ws = wb.create_sheet('Actualizado', wb.sheetnames.index(DIC) + 1)

# ---------- estilos (mismos que el diccionario) ----------
NAVY = 'FF1A1F2E'
navy_fill = PatternFill('solid', fgColor=NAVY)
yellow_fill = PatternFill('solid', fgColor='FFFFFF00')
soft_fill = PatternFill('solid', fgColor='FFF2F4F8')
thin = Side(style='thin', color='FFBFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
h_font = Font(name='Calibri', size=10.5, bold=True, color='FFFFFFFF')
b_font = Font(name='Calibri', size=10)
in_font = Font(name='Calibri', size=11, bold=True, color='FF0000FF')

headers = [src.cell(1, c).value for c in range(1, 15)]

# ---------- encabezado de la hoja ----------
ws['A2'] = 'BUSCADOR DE INDICADORES POR NIVEL'
ws['A2'].font = Font(name='Calibri', size=16, bold=True, color=NAVY)
ws.merge_cells('A2:F2')
ws['A3'] = ("Seleccione un Nivel en la celda amarilla: la tabla trae automáticamente todos los indicadores "
            f"registrados con ese nivel en la hoja '{DIC.strip()}'.")
ws['A3'].font = Font(name='Calibri', size=10, italic=True, color='FF595959')
ws.merge_cells('A3:N3')

ws['A5'] = 'NIVEL:'
ws['A5'].font = Font(name='Calibri', size=11, bold=True, color=NAVY)
ws['A5'].alignment = Alignment(horizontal='right', vertical='center')
ws['B5'] = 'Estratégico'
ws['B5'].font = in_font
ws['B5'].fill = yellow_fill
ws['B5'].border = box
ws['B5'].alignment = Alignment(horizontal='center', vertical='center')

ws['D5'] = 'Indicadores encontrados:'
ws['D5'].font = Font(name='Calibri', size=11, bold=True, color=NAVY)
ws['D5'].alignment = Alignment(horizontal='right', vertical='center')
ws['E5'] = f'=COUNTIF({Q}!$M${FIRST}:$M${LAST},{LVL})'
ws['E5'].font = Font(name='Calibri', size=11, bold=True, color=NAVY)
ws['E5'].fill = soft_fill
ws['E5'].border = box
ws['E5'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[5].height = 20

dv = DataValidation(type='list', formula1="'Listas'!$C$1:$C$3", allow_blank=False,
                    showDropDown=False, showErrorMessage=True,
                    errorTitle='Nivel no válido',
                    error='Elija un nivel de la lista: Estratégico, Táctico u Operativo.',
                    promptTitle='Nivel', prompt='Seleccione el nivel a consultar.')
ws.add_data_validation(dv)
dv.add(ws['B5'])

# ---------- cabecera de la tabla ----------
for i, h in enumerate(headers, start=1):
    c = ws.cell(HDR_ROW, i, h)
    c.font, c.fill, c.border = h_font, navy_fill, box
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[HDR_ROW].height = 33.75

# ---------- fórmulas ----------
for i in range(1, 15):
    L = get_column_letter(i)
    for k, r in enumerate(range(R0, RN + 1), start=1):
        if MODE == 'aggregate':
            f = (f'=IFERROR(INDEX({Q}!{L}${FIRST}:{L}${LAST},'
                 f'AGGREGATE(15,6,(ROW({Q}!$M${FIRST}:$M${LAST})-{FIRST-1})'
                 f'/({Q}!$M${FIRST}:$M${LAST}={LVL}),ROWS($A${R0}:$A{r}))),"")')
        else:
            f = (f'=IFERROR(INDEX({Q}!{L}${FIRST}:{L}${LAST},'
                 f'MATCH(ROWS($A${R0}:$A{r}),$P${R0}:$P${RN},0)),"")')
        c = ws.cell(r, i, f)
        c.font = b_font
        c.border = box
        c.alignment = Alignment(vertical='top', wrap_text=True)
        if (k % 2) == 0:
            c.fill = soft_fill

# columna auxiliar (oculta): ranking de coincidencias
if MODE == 'helper':
    ws['P6'] = 'Auxiliar (no borrar)'
    ws['P6'].font = Font(name='Calibri', size=8, italic=True, color='FF808080')
    for r in range(R0, RN + 1):
        srow = FIRST + (r - R0)
        ws.cell(r, 16, f'=IF({Q}!$M{srow}={LVL},COUNTIF({Q}!$M${FIRST}:$M{srow},{LVL}),"")')
    ws.column_dimensions['P'].hidden = True

# ---------- nota al pie ----------
nota = ws.cell(RN + 2, 1)
if MODE == 'helper':
    nota.value = ('Cómo funciona: la columna auxiliar P (oculta) numera las filas del diccionario que coinciden con el nivel '
                  'elegido en B5, y cada celda de la tabla usa INDEX + MATCH para traer la coincidencia n.º de esa lista. '
                  'Si usa Excel 365 puede reemplazar todo por una sola fórmula dinámica en A8: '
                  f'=FILTER({Q}!A{FIRST}:N{LAST},{Q}!M{FIRST}:M{LAST}={LVL},"Sin resultados")')
else:
    nota.value = ('Cómo funciona: AGGREGATE(15;6;...) devuelve la fila del diccionario que ocupa la posición n entre las que '
                  'coinciden con el nivel elegido en B5, e INDEX trae el dato de esa fila. '
                  'Si usa Excel 365 puede reemplazar todo por una sola fórmula dinámica en A8: '
                  f'=FILTER({Q}!A{FIRST}:N{LAST},{Q}!M{FIRST}:M{LAST}={LVL},"Sin resultados")')
nota.font = Font(name='Calibri', size=9, italic=True, color='FF595959')
ws.merge_cells(start_row=RN + 2, start_column=1, end_row=RN + 2, end_column=14)
ws.cell(RN + 2, 1).alignment = Alignment(vertical='top', wrap_text=True)
ws.row_dimensions[RN + 2].height = 30

nota2 = ws.cell(RN + 4, 1)
nota2.value = (f'Origen de los datos: hoja "{DIC}", filas {FIRST} a {LAST} (57 indicadores). '
               'Si agrega indicadores más abajo, amplíe el rango $58 en las fórmulas.')
nota2.font = Font(name='Calibri', size=9, italic=True, color='FF595959')
ws.merge_cells(start_row=RN + 4, start_column=1, end_row=RN + 4, end_column=14)

# ---------- anchos / vista ----------
widths = {'A': 8, 'B': 16, 'C': 18, 'D': 18, 'E': 34, 'F': 30, 'G': 60, 'H': 55,
          'I': 16, 'J': 12, 'K': 20, 'L': 16, 'M': 12, 'N': 16}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = f'A{R0}'
ws.sheet_view.showGridLines = False
ws.auto_filter.ref = f'A{HDR_ROW}:N{RN}'

wb.save(OUT)
print('OK ->', OUT, '| modo:', MODE)
